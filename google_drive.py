"""Google Drive API integration module."""
import os
import pickle
from datetime import datetime
from typing import List, Dict, Optional
from googleapiclient.discovery import build
from google_auth_oauthlib.flow import InstalledAppFlow
from google.auth.transport.requests import Request
from googleapiclient.http import MediaIoBaseDownload
import io


# Google Drive API のスコープ
SCOPES = ['https://www.googleapis.com/auth/drive.readonly']


class GoogleDriveMonitor:
    """Google Driveを監視するクラス"""

    def __init__(self, credentials_file: str = 'credentials.json'):
        """
        初期化

        Args:
            credentials_file: Google Drive API の認証情報ファイルパス
        """
        self.credentials_file = credentials_file
        self.service = None
        self.last_check_time = None
        self._authenticate()

    def _authenticate(self):
        """Google Drive APIの認証を行う"""
        import base64

        creds = None

        # 環境変数からcredentials.jsonを復元（クラウドデプロイ用）
        credentials_base64 = os.getenv('GOOGLE_CREDENTIALS_BASE64')
        if credentials_base64 and not os.path.exists(self.credentials_file):
            try:
                credentials_json = base64.b64decode(credentials_base64).decode('utf-8')
                with open(self.credentials_file, 'w') as f:
                    f.write(credentials_json)
                print("環境変数からcredentials.jsonを復元しました")
            except Exception as e:
                print(f"credentials.json復元エラー: {e}")

        # 環境変数からtoken.pickleを復元（クラウドデプロイ用）
        token_base64 = os.getenv('GOOGLE_TOKEN_BASE64')
        if token_base64 and not os.path.exists('token.pickle'):
            try:
                token_data = base64.b64decode(token_base64)
                with open('token.pickle', 'wb') as token:
                    token.write(token_data)
                print("環境変数からtoken.pickleを復元しました")
            except Exception as e:
                print(f"token.pickle復元エラー: {e}")

        # token.pickleファイルがあれば読み込む
        if os.path.exists('token.pickle'):
            with open('token.pickle', 'rb') as token:
                creds = pickle.load(token)

        # 認証情報がないか、無効な場合は再認証
        if not creds or not creds.valid:
            if creds and creds.expired and creds.refresh_token:
                creds.refresh(Request())
            else:
                flow = InstalledAppFlow.from_client_secrets_file(
                    self.credentials_file, SCOPES)
                creds = flow.run_local_server(port=0)

            # 認証情報を保存
            with open('token.pickle', 'wb') as token:
                pickle.dump(creds, token)

        self.service = build('drive', 'v3', credentials=creds)
        print("Google Drive API認証成功")

    def get_recent_files(self, folder_id: Optional[str] = None,
                        since: Optional[datetime] = None) -> List[Dict]:
        """
        最近のファイル一覧を取得

        Args:
            folder_id: 監視するフォルダのID（Noneの場合は全体）
            since: この時刻以降に作成されたファイルのみ取得

        Returns:
            ファイル情報のリスト
        """
        try:
            query_parts = []

            # フォルダ指定がある場合
            if folder_id:
                query_parts.append(f"'{folder_id}' in parents")

            # 時刻指定がある場合
            if since:
                time_str = since.strftime('%Y-%m-%dT%H:%M:%S')
                query_parts.append(f"createdTime > '{time_str}'")

            # ゴミ箱以外のファイル
            query_parts.append("trashed = false")

            query = ' and '.join(query_parts) if query_parts else None

            results = self.service.files().list(
                q=query,
                pageSize=100,
                fields="nextPageToken, files(id, name, mimeType, createdTime, modifiedTime, size)",
                orderBy="createdTime desc"
            ).execute()

            files = results.get('files', [])
            return files

        except Exception as e:
            print(f"ファイル一覧取得エラー: {e}")
            return []

    def download_file_content(self, file_id: str, mime_type: str) -> Optional[bytes]:
        """
        ファイルの内容をダウンロード

        Args:
            file_id: ファイルID
            mime_type: MIMEタイプ

        Returns:
            ファイルの内容（バイト列）
        """
        try:
            # Google Docs形式の場合はエクスポート
            if 'google-apps' in mime_type:
                # Google Docsをテキストとしてエクスポート
                if 'document' in mime_type:
                    request = self.service.files().export_media(
                        fileId=file_id,
                        mimeType='text/plain'
                    )
                # Google Sheetsをテキストとしてエクスポート
                elif 'spreadsheet' in mime_type:
                    request = self.service.files().export_media(
                        fileId=file_id,
                        mimeType='text/csv'
                    )
                # その他のGoogle Apps形式はPDFでエクスポート
                else:
                    request = self.service.files().export_media(
                        fileId=file_id,
                        mimeType='application/pdf'
                    )
            else:
                # 通常のファイルはそのままダウンロード
                request = self.service.files().get_media(fileId=file_id)

            file_buffer = io.BytesIO()
            downloader = MediaIoBaseDownload(file_buffer, request)

            done = False
            while not done:
                status, done = downloader.next_chunk()

            return file_buffer.getvalue()

        except Exception as e:
            print(f"ファイルダウンロードエラー ({file_id}): {e}")
            return None

    def extract_text_from_content(self, content: bytes, mime_type: str,
                                  filename: str) -> str:
        """
        ファイル内容からテキストを抽出

        Args:
            content: ファイルの内容（バイト列）
            mime_type: MIMEタイプ
            filename: ファイル名

        Returns:
            抽出されたテキスト
        """
        try:
            # テキストファイル
            if 'text/' in mime_type or mime_type == 'application/json':
                return content.decode('utf-8', errors='ignore')

            # PDF
            elif mime_type == 'application/pdf':
                import PyPDF2
                pdf_file = io.BytesIO(content)
                pdf_reader = PyPDF2.PdfReader(pdf_file)
                text = ""
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
                return text

            # Word文書
            elif mime_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                from docx import Document
                doc_file = io.BytesIO(content)
                doc = Document(doc_file)
                text = "\n".join([para.text for para in doc.paragraphs])
                return text

            # Excel
            elif mime_type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
                from openpyxl import load_workbook
                excel_file = io.BytesIO(content)
                wb = load_workbook(excel_file)
                text = ""
                for sheet in wb.worksheets:
                    text += f"\n=== {sheet.title} ===\n"
                    for row in sheet.iter_rows(values_only=True):
                        text += ", ".join([str(cell) if cell else "" for cell in row]) + "\n"
                return text

            # その他のファイルはバイナリとして扱う
            else:
                return f"[バイナリファイル: {filename}, タイプ: {mime_type}]"

        except Exception as e:
            print(f"テキスト抽出エラー: {e}")
            return f"[テキスト抽出失敗: {filename}]"

    def check_new_files(self, folder_id: Optional[str] = None) -> List[Dict]:
        """
        新しいファイルをチェック

        Args:
            folder_id: 監視するフォルダのID

        Returns:
            新しいファイルのリスト
        """
        # 初回チェックの場合は現在時刻を記録して終了
        if self.last_check_time is None:
            self.last_check_time = datetime.utcnow()
            print(f"監視開始: {self.last_check_time}")
            return []

        # 前回チェック以降のファイルを取得
        new_files = self.get_recent_files(folder_id, self.last_check_time)

        # 最終チェック時刻を更新
        self.last_check_time = datetime.utcnow()

        return new_files
